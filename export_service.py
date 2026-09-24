"""
Portfolio & Proposal Export Service for BT_301.
Course: Introduction to Biotechnology (BT_301).
Generates submission-ready Microsoft Word (.docx) and Markdown (.md) documents
matching the official 10-section university assignment template and rubric.
Authored & Architected by Maya Abdelrazek & Youssef Aboulkheir.
"""

import io
from typing import Dict, Any, List
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _set_cell_background(cell, hex_color: str):
    """Sets background shading of a docx table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)


def _get_team_members(data: Dict[str, Any]) -> List[Dict[str, str]]:
    """Extracts team members list or falls back to primary student info."""
    raw = data.get("team_members", [])
    if isinstance(raw, list) and len(raw) > 0 and any(m.get("name") for m in raw):
        return [
            {
                "name": m.get("name", "Student").strip(),
                "id": m.get("id", "—").strip(),
                "email": m.get("email", "—").strip()
            }
            for m in raw
            if m.get("name") or m.get("id") or m.get("email")
        ]
    # Fallback to single student
    s_name = data.get("student_name", "BT_301 Student / Team")
    s_id = data.get("student_id", "—")
    return [{"name": s_name, "id": s_id, "email": "—"}]


def generate_markdown_portfolio(data: Dict[str, Any], audit_results: Dict[str, Any] = None) -> str:
    """
    Generates a full Markdown proposal aligned 1:1 with the official BT_301 template.
    """
    title = data.get("title", "Untitled Proposal")
    chassis = data.get("chassis", data.get("system", "N/A"))
    tool = data.get("tool", "N/A")
    target = data.get("target", "N/A")
    sub_date = data.get("submission_date") or datetime.now().strftime("%B %d, %Y")
    team = _get_team_members(data)

    md = []
    # Official University Header
    md.append(f"# {title}")
    md.append("## Introduction to Biotechnology (BT_301)")
    md.append("### Project Information\n")
    md.append(f"**Project Title:** {title}\n")
    md.append("#### Team Members")
    md.append("| Name | ID | Email |")
    md.append("| :--- | :--- | :--- |")
    for m in team:
        md.append(f"| {m.get('name', '—')} | {m.get('id', '—')} | {m.get('email', '—')} |")
    md.append(f"\n**Submission Date:** {sub_date}")
    md.append(f"**Target Agency Alignment:** STDF / ICGEB Grant Proposal Standard")
    md.append(f"**Academic Framework & System Architects:** Maya Abdelrazek & Youssef Aboulkheir")
    md.append("\n---\n")

    # 1. Abstract & Keywords
    md.append("# 1. Abstract")
    abstract_text = data.get("abstract", "*No abstract provided.*")
    md.append(abstract_text + "\n")
    kw = data.get("keywords", "")
    if kw:
        md.append(f"**Keywords:** {kw}\n")
    else:
        md.append("**Keywords:** Biotechnology, Molecular Engineering, Chassis, Empirical Benchmark\n")
    md.append("---\n")

    # 2. Introduction & Background
    md.append("# 2. Introduction & Background")
    md.append("## 2.1 Problem Statement & Background")
    narrative = data.get("problem_narrative")
    if not narrative:
        funnel = data.get("funnel", {})
        if any(funnel.values()):
            parts = [funnel.get("tier1", ""), funnel.get("tier2", ""), funnel.get("tier3", ""), funnel.get("tier4", "")]
            narrative = " ".join([p for p in parts if p]).strip()
    if not narrative and data.get("matrix"):
        m = data.get("matrix", {})
        narrative = f"{m.get('yellow', '')} {m.get('blue', '')} {m.get('green', '')} {m.get('red', '')}".strip()
    md.append(narrative if narrative else "*Unified background narrative pending.*")
    md.append("\n")

    md.append("## 2.2 Project Importance & Multi-Domain Impact")
    impact_domains = data.get("impact_domains", [])
    if impact_domains:
        md.append(f"**Key Target Domains:** {', '.join([d.capitalize() for d in impact_domains])}\n")
    impact_text = data.get("impact_text") or data.get("impact")
    md.append(impact_text if impact_text else "*Multi-domain translational impact pending.*")
    md.append("\n")

    md.append("## 2.3 Target Audience & Direct Beneficiaries")
    aud = data.get("target_audience") or data.get("customer")
    md.append(aud if aud else "*Target beneficiaries and end-users pending.*")
    md.append("\n---\n")

    # 3. Project Aim & Objectives
    md.append("# 3. Project Aim & Objectives")
    md.append("## 3.1 Overall Aim")
    aim_ov = data.get("overarching_aim") or data.get("aim")
    if not aim_ov and data.get("funnel", {}).get("tier4"):
        aim_ov = data["funnel"]["tier4"]
    md.append(aim_ov if aim_ov else "*Overarching project aim pending.*")
    md.append("\n")

    md.append("## 3.2 Specific Objectives")
    objectives = data.get("objectives", [])
    if objectives and isinstance(objectives, list):
        for idx, obj in enumerate(objectives, 1):
            v = obj.get("verb", "Execute")
            t = obj.get("text", "")
            md.append(f"{idx}. **{v}** {t}")
    else:
        aims = data.get("aims", {})
        for i in range(1, 4):
            val = aims.get(f"aim{i}")
            if val:
                title_s = f" ({aims.get(f'aim{i}_title')})" if aims.get(f"aim{i}_title") else ""
                md.append(f"{i}. **Specific Aim {i}{title_s}:** {val}")
    md.append("\n---\n")

    # 4. Methodology & Experimental Design
    md.append("# 4. Methodology & Experimental Design")
    md.append("## 4.1 Technical Overview & Workflow")
    md.append(data.get("methodology", "*Methodological workflow overview pending (Future Tense).*") + "\n")

    md.append("## 4.2 Workflow Diagram")
    fc_cap = data.get("flowchart_caption", "Figure 1: Schematic workflow depicting genetic engineering and functional validation pipeline.")
    md.append(f"*{fc_cap}*\n")

    md.append("## 4.3 Materials & Methods")
    methods_text = data.get("materials_methods") or ""
    if methods_text:
        md.append(methods_text + "\n")
    controls = data.get("controls_checklist", {})
    md.append("### Rigorous Experimental Controls & Triplicates Standard")
    md.append(f"- **Positive Control**: {'[x] Calibrated reference included' if controls.get('positive') else '[ ] Reference standard'}")
    md.append(f"- **Negative / Mock Control**: {'[x] Empty vector / uninduced host included' if controls.get('negative') else '[ ] Baseline control'}")
    md.append(f"- **Biological & Technical Triplicates (n=3)**: {'[x] Triplicate statistical reproducibility enabled' if controls.get('triplicates') else '[ ] Triplicate design'}\n")
    md.append("---\n")

    # 5. Risk Management & Contingency Plan
    md.append("# 5. Risk Management & Contingency Plan")
    risks = data.get("risks", [])
    if risks and isinstance(risks, list):
        md.append("| Potential Technical Failure Mode | Risk Level | Fallback Contingency Plan (Plan B) |")
        md.append("| :--- | :---: | :--- |")
        for r in risks:
            md.append(f"| {r.get('risk', '—')} | {r.get('level', 'Moderate')} | {r.get('fallback', '—')} |")
    else:
        cont = data.get("aims_contingencies", {})
        md.append("| Specific Objective | Potential Pitfall | Fallback Contingency Alternative |")
        md.append("| :--- | :--- | :--- |")
        md.append(f"| **Objective 1** | Suboptimal construct expression | {cont.get('aim1_fallback', 'Codon optimize and switch to alternative promoter system')} |")
        md.append(f"| **Objective 2** | Low activity / misfolding | {cont.get('aim2_fallback', 'Introduce stabilizing salt bridges or chaperone co-expression')} |")
        md.append(f"| **Objective 3** | Matrix inhibition in real specimen | {cont.get('aim3_fallback', 'Incorporate sample pre-treatment and filtration step')} |")
    md.append("\n---\n")

    # 6. Ansoff Matrix for Biotechnology
    md.append("# 6. Ansoff Matrix for Biotechnology")
    ansoff = data.get("ansoff", {})
    md.append("| Strategic Quadrant | Focus & Definition | Application to Proposed Research |")
    md.append("| :--- | :--- | :--- |")
    md.append(f"| **Market Penetration** | Existing Technology / Existing Market | {ansoff.get('penetration', 'Optimizing catalytic kinetics and bioprocess throughput')} |")
    md.append(f"| **Product Development** | New Engineered Modality / Existing Market | {ansoff.get('prod_dev', 'Deploying engineered novel enzyme with elevated thermostability')} |")
    md.append(f"| **Market Development** | Existing Technology / New Market | {ansoff.get('mkt_dev', 'Translating wastewater biocatalyst to agricultural effluents')} |")
    md.append(f"| **Diversification** | New Technology / New Market | {ansoff.get('diversification', 'Pioneering living self-healing biomaterials in municipal infrastructure')} |")
    md.append("\n---\n")

    # 7. SWOT Analysis
    md.append("# 7. SWOT Analysis")
    swot = data.get("swot", {})
    md.append("| Strengths (Internal) | Weaknesses (Internal) |")
    md.append("| :--- | :--- |")
    md.append(f"| {swot.get('strengths', 'High catalytic specificity and non-toxic chassis')} | {swot.get('weaknesses', 'Potential recombinant expression bottlenecks')} |")
    md.append("| **Opportunities (External)** | **Threats (External)** |")
    md.append(f"| {swot.get('opportunities', 'Surging demand for sustainable bio-solutions')} | {swot.get('threats', 'Regulatory approval timelines for GMO derivatives')} |")
    md.append("\n---\n")

    # 8. PESTEL Analysis
    md.append("# 8. PESTEL Analysis")
    pestel = data.get("pestel", {})
    md.append(f"- **Political:** {pestel.get('political', 'National biotechnology investment initiatives and GMO regulatory oversight')}")
    md.append(f"- **Economic:** {pestel.get('economic', 'Cost competitiveness against petro-chemical synthesis in EGP')}")
    md.append(f"- **Social:** {pestel.get('social', 'Public acceptance of engineered biological products and consumer awareness')}")
    md.append(f"- **Technological:** {pestel.get('tech', 'Integration with state-of-the-art synthetic biology and automated screening')}")
    md.append(f"- **Environmental:** {pestel.get('env', 'Zero-toxic byproduct generation and circular economy alignment')}")
    md.append(f"- **Legal:** {pestel.get('legal', 'Compliance with Egyptian biosafety regulations and Cartagena protocol')}\n")
    md.append("---\n")

    # 9. Timeline & Budget
    md.append("# 9. Timeline & Budget")
    md.append("## 9.1 Project Timeline & Milestones")
    gs = data.get("gantt_schedule", {})
    tot_m = gs.get("total_months", 18)
    md.append(f"**Total Staged Duration:** {tot_m} Months (Pure Scientific Milestones - No Student Names Assigned)\n")
    wp1_g = f" | Go/No-Go Gate: {gs.get('wp1_gate')}" if gs.get('wp1_gate') else ""
    wp2_g = f" | Go/No-Go Gate: {gs.get('wp2_gate')}" if gs.get('wp2_gate') else ""
    wp3_g = f" | Go/No-Go Gate: {gs.get('wp3_gate')}" if gs.get('wp3_gate') else ""
    md.append(f"- **Work Package 1 (Design & Molecular Construction):** Months {gs.get('wp1_start', 1)}–{gs.get('wp1_end', 6)}{wp1_g}")
    md.append(f"- **Work Package 2 (Expression & Functional Characterization):** Months {gs.get('wp2_start', 5)}–{gs.get('wp2_end', 12)}{wp2_g}")
    md.append(f"- **Work Package 3 (Optimization & Empirical Benchmarking):** Months {gs.get('wp3_start', 10)}–{gs.get('wp3_end', tot_m)}{wp3_g}\n")

    md.append("## 9.2 Budget & Resource Allocation (EGP)")
    egp_items = data.get("budget_egp_items", [])
    if egp_items and isinstance(egp_items, list):
        md.append("| Item Description & Specs | Official Category | Qty | Unit Cost (EGP) | Total Cost (EGP) |")
        md.append("| :--- | :--- | :---: | ---: | ---: |")
        grand_total_egp = 0.0
        for item in egp_items:
            q = int(item.get("qty", 1))
            u = float(item.get("unit_cost", 0))
            tot = q * u
            grand_total_egp += tot
            md.append(f"| {item.get('desc', 'Item')} | {item.get('category', 'B. Consumables & Reagents')} | {q} | {u:,.2f} EGP | {tot:,.2f} EGP |")
        md.append(f"| **TOTAL ESTIMATED BUDGET** | — | — | — | **{grand_total_egp:,.2f} EGP (ج.م)** |\n")
    else:
        # Standard default EGP table
        md.append("| Official Category | Allocated Subtotal (EGP) | Justification & Deliverable |")
        md.append("| :--- | ---: | :--- |")
        md.append("| A. Equipment & Facilities | 45,000.00 EGP | Core facility bench fees, HPLC & spectrophotometer access |")
        md.append("| B. Consumables & Reagents | 83,000.00 EGP | Primers, gene synthesis, polymerases, MasterMixes, media |")
        md.append("| C. Travel & Field Work | 8,000.00 EGP | Specimen collection, field transport, and cold-chain storage |")
        md.append("| D. Personnel & Student Stipends | 36,000.00 EGP | Student research assistant stipends (3 students x 12,000 EGP) |")
        md.append("| E. Training & Computational/Software | 15,000.00 EGP | Bioinformatics cloud compute, AlphaFold/docking licenses |")
        md.append("| F. Other Direct Costs & Contingency | 15,000.00 EGP | Unforeseen experimental buffer and biohazard waste disposal |")
        md.append("| **TOTAL ESTIMATED BUDGET** | **202,000.00 EGP (ج.م)** | Standard 18-month undergraduate research grant |")
    md.append("\n---\n")

    # 10. References
    md.append("# 10. References")
    refs = data.get("references", "")
    md.append(refs if refs else "*No references entered.*")
    md.append("\n---\n")

    # Proof of Process & Literature Audit Trail (Maintains full compatibility with test_markdown_export)
    md.append("## Proof-of-Process & Literature Audit Trail")
    md.append("### 1.1 The 3 Magic Nouns (Topic Formulation)")
    md.append(f"- **Biological Chassis**: `{chassis}`")
    md.append(f"- **Molecular Tool / Enzyme**: `{tool}`")
    md.append(f"- **Target Problem / Villain**: `{target}`\n")

    md.append("### 1.2 Boolean Search Query Used")
    search_q = data.get("search_query", f'"{chassis}" AND "{tool}" AND "{target}"')
    md.append(f"```text\n{search_q}\n```\n")

    md.append("### 1.3 4-Color Evidence Extraction Matrix")
    matrix = data.get("matrix", {})
    md.append("| Color / Category | Extracted Primary Literature Evidence | Metric / Baseline |")
    md.append("| :--- | :--- | :--- |")
    md.append(f"| **🟡 Yellow: Crisis Burden** | {matrix.get('yellow', 'N/A')} | Hard number/prevalence |")
    md.append(f"| **🔵 Blue: Current Tool** | {matrix.get('blue', 'N/A')} | Existing baseline |")
    md.append(f"| **🟢 Green: Knowledge Gap** | {matrix.get('green', 'N/A')} | Critical bottleneck |")
    md.append(f"| **🔴 Red: Target Milestone** | {matrix.get('red', 'N/A')} | Target quantitative threshold |")
    md.append("\n---\n")

    # Section 3: Socratic Rubric Audit
    if audit_results:
        md.append("## Faculty Socratic Audit & Evaluation Report")
        raw = audit_results.get("raw_total", 0)
        gpa = audit_results.get("gpa_score", 0)
        md.append(f"**Total Raw Score**: `{raw} / 85.0 marks`  ")
        md.append(f"**Estimated BT_301 GPA Grade**: `{gpa} / 10.0 GPA`\n")

    md.append("\n---\n")
    md.append("### Protected Intellectual Property & Academic Framework")
    md.append("© BioWriter Studio • Research Proposal & Proof-of-Process Learning Architecture  ")
    md.append("Curated & Architected by **Maya Abdelrazek & Youssef Aboulkheir**. All rights reserved.\n")

    return "\n".join(md)


def generate_docx_portfolio(data: Dict[str, Any], audit_results: Dict[str, Any] = None) -> io.BytesIO:
    """
    Builds a professional Microsoft Word (.docx) portfolio matching the BT_301 official template.
    """
    doc = Document()

    # Core Properties
    doc.core_properties.author = "Maya Abdelrazek & Youssef Aboulkheir"
    doc.core_properties.comments = "BioWriter Studio Academic Framework. Curated by Maya Abdelrazek & Youssef Aboulkheir."

    # Page Margins: 1 inch & Footer Watermark
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

        footer = section.footer
        footer_p = footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer_run = footer_p.add_run("BioWriter Studio • BT_301 Proposal Standard | Maya Abdelrazek & Youssef Aboulkheir")
        footer_run.font.size = Pt(8.5)
        footer_run.font.italic = True
        footer_run.font.color.rgb = RGBColor(148, 163, 184)

    # University Course Header
    course_p = doc.add_paragraph()
    c_run = course_p.add_run("Introduction to Biotechnology\n(BT_301)")
    c_run.font.size = Pt(14)
    c_run.font.bold = True
    c_run.font.color.rgb = RGBColor(67, 56, 202) # Indigo
    course_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    course_p.paragraph_format.space_after = Pt(12)

    # Project Information Heading
    doc.add_heading("Project Information", level=1)
    title = data.get("title", "Untitled Research Grant Proposal")
    p_title = doc.add_paragraph()
    p_title.add_run("Project Title:\n").bold = True
    r_tval = p_title.add_run(title)
    r_tval.font.size = Pt(12)
    r_tval.font.bold = True
    p_title.paragraph_format.space_after = Pt(10)

    # Team Members Table
    doc.add_heading("Team Members", level=2)
    team = _get_team_members(data)
    team_table = doc.add_table(rows=len(team) + 1, cols=3)
    team_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_headers = ["Name", "ID", "Email"]
    for idx, h in enumerate(t_headers):
        cell = team_table.rows[0].cells[idx]
        cell.text = h
        _set_cell_background(cell, "1E293B")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        cell.paragraphs[0].runs[0].font.bold = True

    for row_idx, member in enumerate(team, 1):
        r = team_table.rows[row_idx]
        r.cells[0].text = member.get("name", "—")
        r.cells[1].text = member.get("id", "—")
        r.cells[2].text = member.get("email", "—")
        if row_idx % 2 == 0:
            for c_idx in range(3):
                _set_cell_background(r.cells[c_idx], "F8FAFC")

    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    sub_date = data.get("submission_date") or datetime.now().strftime("%B %d, %Y")
    p_date = doc.add_paragraph()
    p_date.add_run("Submission Date: ").bold = True
    p_date.add_run(sub_date)
    p_date.paragraph_format.space_after = Pt(16)

    # 1. Abstract
    doc.add_heading("1. Abstract", level=1)
    ab_p = doc.add_paragraph()
    ab_p.add_run(data.get("abstract", "No structured abstract provided."))
    ab_p.paragraph_format.space_after = Pt(8)
    kw = data.get("keywords", "Biotechnology, Molecular Engineering, Chassis, Empirical Benchmark")
    p_kw = doc.add_paragraph()
    p_kw.add_run("Keywords: ").bold = True
    p_kw.add_run(kw)
    p_kw.paragraph_format.space_after = Pt(14)

    # 2. Introduction & Background
    doc.add_heading("2. Introduction & Background", level=1)
    doc.add_heading("2.1 Problem Statement & Background", level=2)
    narrative = data.get("problem_narrative")
    if not narrative:
        funnel = data.get("funnel", {})
        if any(funnel.values()):
            parts = [funnel.get("tier1", ""), funnel.get("tier2", ""), funnel.get("tier3", ""), funnel.get("tier4", "")]
            narrative = " ".join([p for p in parts if p]).strip()
    if not narrative and data.get("matrix"):
        m = data.get("matrix", {})
        narrative = f"{m.get('yellow', '')} {m.get('blue', '')} {m.get('green', '')} {m.get('red', '')}".strip()
    p_bg = doc.add_paragraph()
    p_bg.add_run(narrative if narrative else "[Unified background narrative pending]")
    p_bg.paragraph_format.space_after = Pt(10)

    doc.add_heading("2.2 Project Importance & Multi-Domain Impact", level=2)
    p_imp = doc.add_paragraph()
    impact_domains = data.get("impact_domains", [])
    if impact_domains:
        p_imp.add_run(f"Target Domains: {', '.join([d.capitalize() for d in impact_domains])}\n").bold = True
    p_imp.add_run(data.get("impact_text") or data.get("impact") or "[Multi-domain impact statement pending]")
    p_imp.paragraph_format.space_after = Pt(10)

    doc.add_heading("2.3 Target Audience & Direct Beneficiaries", level=2)
    p_aud = doc.add_paragraph()
    p_aud.add_run(data.get("target_audience") or data.get("customer") or "[Target beneficiaries pending]")
    p_aud.paragraph_format.space_after = Pt(14)

    # 3. Project Aim & Objectives
    doc.add_heading("3. Project Aim & Objectives", level=1)
    doc.add_heading("3.1 Overall Aim", level=2)
    aim_ov = data.get("overarching_aim") or data.get("aim") or data.get("funnel", {}).get("tier4") or "[Overall project aim pending]"
    p_aim = doc.add_paragraph()
    p_aim.add_run(aim_ov)
    p_aim.paragraph_format.space_after = Pt(10)

    doc.add_heading("3.2 Specific Objectives", level=2)
    objectives = data.get("objectives", [])
    if objectives and isinstance(objectives, list):
        for idx, obj in enumerate(objectives, 1):
            ap = doc.add_paragraph(style='List Bullet')
            ap.add_run(f"Objective {idx} ({obj.get('verb', 'Execute')}): ").bold = True
            ap.add_run(obj.get("text", ""))
    else:
        aims = data.get("aims", {})
        for i in range(1, 4):
            val = aims.get(f"aim{i}")
            if val:
                ap = doc.add_paragraph(style='List Bullet')
                ap.add_run(f"Specific Aim {i}: ").bold = True
                ap.add_run(val)
    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 4. Methodology & Experimental Design
    doc.add_heading("4. Methodology & Experimental Design", level=1)
    doc.add_heading("4.1 Technical Overview & Workflow (Future Tense)", level=2)
    p_m1 = doc.add_paragraph()
    p_m1.add_run(data.get("methodology", "[Technical overview pending]"))
    p_m1.paragraph_format.space_after = Pt(10)

    doc.add_heading("4.2 Workflow Diagram", level=2)
    fc_p = doc.add_paragraph()
    fc_p.add_run(data.get("flowchart_caption", "Figure 1: Methodological Workflow Pipeline")).italic = True
    fc_p.paragraph_format.space_after = Pt(10)

    doc.add_heading("4.3 Materials & Methods", level=2)
    p_m3 = doc.add_paragraph()
    p_m3.add_run(data.get("materials_methods") or "[Materials, strains, and protocols pending]")
    p_m3.paragraph_format.space_after = Pt(8)

    # Controls Checklist
    controls = data.get("controls_checklist", {})
    p_ctrl = doc.add_paragraph()
    p_ctrl.add_run("Experimental Controls & Quality Assurance:\n").bold = True
    p_ctrl.add_run(f"• Positive Control: {'[VERIFIED] Benchmark reference included' if controls.get('positive') else '[PENDING] Reference standard'}\n")
    p_ctrl.add_run(f"• Negative Control: {'[VERIFIED] Mock/uninduced vehicle included' if controls.get('negative') else '[PENDING] Vehicle control'}\n")
    p_ctrl.add_run(f"• Replicates: {'[VERIFIED] Triplicates (n=3) with standard deviation' if controls.get('triplicates') else '[PENDING] Biological triplicates'}\n")
    p_ctrl.paragraph_format.space_after = Pt(14)

    # 5. Risk Management & Contingency Plan
    doc.add_heading("5. Risk Management & Contingency Plan", level=1)
    risks = data.get("risks", [])
    if risks and isinstance(risks, list):
        r_table = doc.add_table(rows=len(risks) + 1, cols=3)
        r_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for idx, h in enumerate(["Failure Mode", "Severity / Likelihood", "Fallback Plan B"]):
            c = r_table.rows[0].cells[idx]
            c.text = h
            _set_cell_background(c, "1E293B")
            c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            c.paragraphs[0].runs[0].font.bold = True
        for idx, r in enumerate(risks, 1):
            row = r_table.rows[idx]
            row.cells[0].text = str(r.get("risk", "—"))
            row.cells[1].text = str(r.get("level", "Moderate"))
            row.cells[2].text = str(r.get("fallback", "—"))
            if idx % 2 == 0:
                for c_idx in range(3):
                    _set_cell_background(row.cells[c_idx], "F8FAFC")
    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # 6. Ansoff Matrix for Biotechnology
    doc.add_heading("6. Ansoff Matrix for Biotechnology", level=1)
    ansoff = data.get("ansoff", {})
    ansoff_table = doc.add_table(rows=5, cols=3)
    ansoff_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, h in enumerate(["Quadrant", "Strategic Focus", "Proposal Application"]):
        c = ansoff_table.rows[0].cells[idx]
        c.text = h
        _set_cell_background(c, "1E293B")
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        c.paragraphs[0].runs[0].font.bold = True

    a_rows = [
        ("Market Penetration", "Existing Product / Existing Market", ansoff.get("penetration", "Catalytic optimization and yield elevation")),
        ("Product Development", "New Product / Existing Market", ansoff.get("prod_dev", "Engineered novel construct with enhanced stability")),
        ("Market Development", "Existing Product / New Market", ansoff.get("mkt_dev", "Translating wastewater bioprocess to agricultural runoffs")),
        ("Diversification", "New Product / New Market", ansoff.get("diversification", "Living self-repairing synthetic materials"))
    ]
    for idx, (q, f, app) in enumerate(a_rows, 1):
        r = ansoff_table.rows[idx]
        r.cells[0].text = q
        r.cells[1].text = f
        r.cells[2].text = app
        if idx % 2 == 0:
            for c_idx in range(3):
                _set_cell_background(r.cells[c_idx], "F8FAFC")
    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # 7. SWOT Analysis
    doc.add_heading("7. SWOT Analysis", level=1)
    swot = data.get("swot", {})
    swot_table = doc.add_table(rows=3, cols=2)
    swot_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    swot_table.rows[0].cells[0].text = "Internal Factors"
    swot_table.rows[0].cells[1].text = "External Factors"
    _set_cell_background(swot_table.rows[0].cells[0], "1E293B")
    _set_cell_background(swot_table.rows[0].cells[1], "1E293B")
    swot_table.rows[0].cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    swot_table.rows[0].cells[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    swot_table.rows[1].cells[0].text = f"Strengths:\n{swot.get('strengths', 'High catalytic efficiency')}"
    swot_table.rows[1].cells[1].text = f"Opportunities:\n{swot.get('opportunities', 'Growing market for green biocatalysts')}"
    swot_table.rows[2].cells[0].text = f"Weaknesses:\n{swot.get('weaknesses', 'Potential expression bottlenecks')}"
    swot_table.rows[2].cells[1].text = f"Threats:\n{swot.get('threats', 'Regulatory approval timelines')}"
    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # 8. PESTEL Analysis
    doc.add_heading("8. PESTEL Analysis", level=1)
    pestel = data.get("pestel", {})
    for dim, val in [
        ("Political", pestel.get("political", "National biotechnology support policies")),
        ("Economic", pestel.get("economic", "Cost savings in Egyptian Pounds (EGP)")),
        ("Social", pestel.get("social", "Public awareness and bioethics acceptance")),
        ("Technological", pestel.get("tech", "Rapid advances in gene editing tools")),
        ("Environmental", pestel.get("env", "Circular economy and zero-toxic effluents")),
        ("Legal", pestel.get("legal", "Egyptian biosafety guidelines compliance"))
    ]:
        p_dim = doc.add_paragraph(style='List Bullet')
        p_dim.add_run(f"{dim}: ").bold = True
        p_dim.add_run(val)
    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # 9. Timeline & Budget
    doc.add_heading("9. Timeline & Budget", level=1)
    doc.add_heading("9.1 Project Timeline & Milestones", level=2)
    gs = data.get("gantt_schedule", {})
    tot_m = gs.get("total_months", 18)
    p_time = doc.add_paragraph()
    p_time.add_run(f"Total Staged Project Duration: {tot_m} Months (Scientific Milestones — Zero Student Names Assigned)\n").bold = True
    p_time.add_run(f"• WP1: Design & Construction (Months {gs.get('wp1_start', 1)}–{gs.get('wp1_end', 6)}) | Gate: {gs.get('wp1_gate', 'Sequence verified')}\n")
    p_time.add_run(f"• WP2: Functional Characterization (Months {gs.get('wp2_start', 5)}–{gs.get('wp2_end', 12)}) | Gate: {gs.get('wp2_gate', 'Activity threshold met')}\n")
    p_time.add_run(f"• WP3: Validation & Optimization (Months {gs.get('wp3_start', 10)}–{gs.get('wp3_end', tot_m)}) | Gate: {gs.get('wp3_gate', 'Benchmark outperformed')}\n")
    p_time.paragraph_format.space_after = Pt(10)

    doc.add_heading("9.2 Budget & Resource Allocation (Egyptian Pounds - EGP)", level=2)
    egp_items = data.get("budget_egp_items", [])
    if egp_items and isinstance(egp_items, list):
        b_table = doc.add_table(rows=len(egp_items) + 2, cols=5)
        b_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for idx, h in enumerate(["Item Description", "Category", "Qty", "Unit (EGP)", "Total (EGP)"]):
            c = b_table.rows[0].cells[idx]
            c.text = h
            _set_cell_background(c, "1E293B")
            c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            c.paragraphs[0].runs[0].font.bold = True
        grand_total_egp = 0.0
        for idx, it in enumerate(egp_items, 1):
            row = b_table.rows[idx]
            q = int(it.get("qty", 1))
            u = float(it.get("unit_cost", 0))
            sub = q * u
            grand_total_egp += sub
            row.cells[0].text = str(it.get("desc", "Item"))
            row.cells[1].text = str(it.get("category", "Consumables"))
            row.cells[2].text = str(q)
            row.cells[3].text = f"{u:,.2f}"
            row.cells[4].text = f"{sub:,.2f}"
            if idx % 2 == 0:
                for c_idx in range(5):
                    _set_cell_background(row.cells[c_idx], "F8FAFC")
        # Grand total
        tot_row = b_table.rows[-1]
        tot_row.cells[0].text = "TOTAL ESTIMATED BUDGET"
        tot_row.cells[4].text = f"{grand_total_egp:,.2f} EGP"
        _set_cell_background(tot_row.cells[0], "FEF3C7")
        _set_cell_background(tot_row.cells[4], "FEF3C7")
        tot_row.cells[0].paragraphs[0].runs[0].font.bold = True
        tot_row.cells[4].paragraphs[0].runs[0].font.bold = True
    else:
        # Default EGP Categories Table
        b_table = doc.add_table(rows=8, cols=3)
        b_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for idx, h in enumerate(["Official Category", "Allocated Subtotal (EGP)", "Justification & Deliverables"]):
            c = b_table.rows[0].cells[idx]
            c.text = h
            _set_cell_background(c, "1E293B")
            c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            c.paragraphs[0].runs[0].font.bold = True

        cats_egp = [
            ("A. Equipment & Facilities", "45,000.00 EGP", "Core facility bench fees, HPLC & spectrophotometry"),
            ("B. Consumables & Reagents", "83,000.00 EGP", "Primers, gene synthesis, polymerases, MasterMixes"),
            ("C. Travel & Field Work", "8,000.00 EGP", "Specimen collection and transport"),
            ("D. Personnel & Student Stipends", "36,000.00 EGP", "Student research stipends (3 students x 12,000 EGP)"),
            ("E. Training & Computational/Software", "15,000.00 EGP", "Bioinformatics cloud compute and structural licenses"),
            ("F. Other Direct Costs & Contingency", "15,000.00 EGP", "Experimental contingency and biohazard disposal")
        ]
        for idx, (cat_name, cat_amt, cat_just) in enumerate(cats_egp, 1):
            r = b_table.rows[idx]
            r.cells[0].text = cat_name
            r.cells[1].text = cat_amt
            r.cells[2].text = cat_just
            if idx % 2 == 0:
                for c_idx in range(3):
                    _set_cell_background(r.cells[c_idx], "F8FAFC")
        tot_row = b_table.rows[7]
        tot_row.cells[0].text = "TOTAL ESTIMATED BUDGET"
        tot_row.cells[1].text = "202,000.00 EGP"
        tot_row.cells[2].text = "Standard 18-month undergraduate research grant"
        for c_idx in range(3):
            _set_cell_background(tot_row.cells[c_idx], "FEF3C7")
            if len(tot_row.cells[c_idx].paragraphs[0].runs) > 0:
                tot_row.cells[c_idx].paragraphs[0].runs[0].font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # 10. References
    doc.add_heading("10. References", level=1)
    rp = doc.add_paragraph()
    rp.add_run(data.get("references", "[No references submitted]"))
    rp.paragraph_format.space_after = Pt(14)

    # Audit Scorecard Appendix
    if audit_results:
        doc.add_page_break()
        doc.add_heading("Faculty Socratic Audit Scorecard", level=1)
        raw = audit_results.get("raw_total", 0)
        gpa = audit_results.get("gpa_score", 0)

        score_p = doc.add_paragraph()
        s_run = score_p.add_run(f"Cumulative Score: {raw} / 85.0 Raw Marks  ➔  Estimated GPA: {gpa} / 10.0\n")
        s_run.bold = True
        s_run.font.size = Pt(13)
        s_run.font.color.rgb = RGBColor(197, 48, 48)

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output
